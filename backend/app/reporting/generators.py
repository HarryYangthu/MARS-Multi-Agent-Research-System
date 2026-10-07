"""Editable Office exports of saved reports and observed experiment data."""
from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any

import mistune
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from app.reporting.presentation import write_deck


def report_blocks(pack: dict[str, Any]) -> list[dict[str, Any]]:
    result = mistune.create_markdown(renderer="ast", plugins=["table"])(str(pack.get("report_markdown") or pack.get("report_markdown_excerpt") or ""))
    return result if isinstance(result, list) else []


def plain(token: dict[str, Any]) -> str:
    if token.get("type") in {"softbreak", "linebreak"}:
        return "\n"
    value = str(token.get("raw", "")) + "".join(plain(child) for child in token.get("children", []))
    if token.get("type") in {"link", "image"}:
        return value + f" ({token.get('attrs', {}).get('url', '')})"
    return value


def table_rows(token: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    for section in token.get("children", []):
        if section.get("type") == "table_head":
            rows.append([plain(cell) for cell in section.get("children", [])])
        else:
            rows.extend([[plain(cell) for cell in row.get("children", [])] for row in section.get("children", [])])
    return rows


def _cell(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return json.dumps(value, ensure_ascii=False, allow_nan=False) if isinstance(value, (dict, list)) else value


def write_results_workbook(path: Path, data_pack: dict[str, Any]) -> None:
    book = Workbook()
    del book["Sheet"]
    metrics = data_pack.get("metrics", [])
    keys = list(dict.fromkeys(str(key) for row in metrics for key in row))
    sheets: list[tuple[str, list[list[Any]]]] = [
        ("Overview", [["Field", "Value"], ["Task", data_pack.get("task")], ["Run", data_pack.get("run_id")],
            ["Project", data_pack.get("project")], ["Approved report", data_pack.get("report_approved")],
            ["Report SHA-256", data_pack.get("report_source_sha256")], ["Experiments", len(metrics)]]),
        ("Metrics", [keys or ["No metrics available"], *[[row.get(key) for key in keys] for row in metrics]]),
        ("Sources", [["Source", "SHA-256"], *[[ref, data_pack.get("source_hashes", {}).get(ref)] for ref in data_pack.get("source_refs", [])]]),
        ("Limitations", [["Input limitations"], *[[reason] for reason in data_pack.get("degraded_reasons", [])],
            ["Only recorded observations are exported. Scientific conclusions and limitations are in the approved report."]]),
        ("Report", [["Full approved report"], *[[line] for line in str(data_pack.get("report_markdown", "")).splitlines()]]),
    ]
    for index, curve in enumerate(data_pack.get("curves", []), 1):
        points = [point for point in curve.get("points", []) if isinstance(point, dict)]
        headers = list(dict.fromkeys(str(key) for point in points for key in point))
        sheets.append((f"Steps_{index}", [headers or ["No step records"], *[[point.get(key) for key in headers] for point in points]]))
    for name, rows in sheets:
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append([_cell(value) for value in row])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Imported report text must never become a formula.
        sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="18243A")
        for column in sheet.columns:
            header = str(column[0].value or "")
            sheet.column_dimensions[get_column_letter(column[0].column)].width = 95 if name == "Report" else min(50, max(18, len(header) + 3))
            for cell in column[1:]:
                cell.font = Font(name="Arial", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str))
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000" if "loss" in header or header == "lr" else "0.00"
        curve_key = next((key for key in ("training_loss", "loss", "value") if key in rows[0]), None)
        if name.startswith("Steps_") and curve_key and sheet.max_row > 1:
            chart = LineChart()
            chart.title = f"{data_pack['curves'][int(name.split('_')[1]) - 1].get('experiment_id', name)} · training loss"
            step_key = "optimizer_step" if "optimizer_step" in rows[0] else "sample_index"
            chart.y_axis.title, chart.x_axis.title = f"Recorded {curve_key}", "Optimizer step" if step_key == "optimizer_step" else "Sample index"
            chart.add_data(Reference(sheet, min_col=rows[0].index(curve_key) + 1, min_row=1, max_row=sheet.max_row), titles_from_data=True)
            if step_key in rows[0]:
                chart.set_categories(Reference(sheet, min_col=rows[0].index(step_key) + 1, min_row=2, max_row=sheet.max_row))
            chart.width, chart.height = 24, 12
            sheet.add_chart(chart, f"{get_column_letter(sheet.max_column + 2)}2")
    book.save(path)


def _native_table(document: Any, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=width)
    table.style, table.autofit = "Table Grid", False
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D8DEE9")):
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, cell in enumerate(cells):
            cell.width, cell.text = Inches(6.5 / width), values[column] if column < len(values) else ""
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                for run in paragraph.runs:
                    run.font.size, run.bold = Pt(9), index == 0
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E9EDF2")
                cell._tc.get_or_add_tcPr().append(shade)
        if index == 0:
            table.rows[index]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.add_paragraph()


def write_research_docx(path: Path, data_pack: dict[str, Any]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name, style.font.color.rgb = "Arial", RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    document.core_properties.title, document.core_properties.author = str(data_pack.get("task", "Research report")), "MARS"

    def add(tokens: list[dict[str, Any]], *, list_depth: int = 0) -> None:
        for token in tokens:
            kind = token.get("type")
            if kind == "heading":
                level = min(3, int(token.get("attrs", {}).get("level", 1)))
                document.add_heading(plain(token), level=0 if not document.paragraphs and level == 1 else level)
            elif kind in {"paragraph", "block_text"}:
                document.add_paragraph(plain(token), style="List Bullet" if list_depth else None).paragraph_format.widow_control = True
            elif kind == "table":
                _native_table(document, table_rows(token))
            elif kind == "block_code":
                for run in document.add_paragraph(str(token.get("raw", ""))).runs:
                    run.font.name, run.font.size = "Courier New", Pt(8)
            elif kind == "list":
                add(token.get("children", []), list_depth=list_depth + 1)
            elif token.get("children"):
                add(token["children"], list_depth=list_depth)
    add(report_blocks(data_pack))
    document.add_heading("导出来源", level=1)
    document.add_paragraph(f"任务：{data_pack.get('run_id')}\n报告：{data_pack.get('report_source')}\nSHA-256：{data_pack.get('report_source_sha256')}")
    for plot in data_pack.get("plots", []):
        root = Path(str(data_pack.get("run_root", ""))).resolve()
        image = (root / str(plot["path"])).resolve()
        if image.is_relative_to(root) and image.suffix.lower() in {".png", ".jpg", ".jpeg"}:
            document.add_heading("已记录的训练曲线", level=2)
            document.add_picture(str(image), width=Inches(6.3))
            document.add_paragraph(str(plot["path"]))
    document.save(str(path))


def write_research_deck(path: Path, data_pack: dict[str, Any]) -> None:
    write_deck(path, data_pack, report_blocks(data_pack), plain, table_rows)


def pretty_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False, default=str)
