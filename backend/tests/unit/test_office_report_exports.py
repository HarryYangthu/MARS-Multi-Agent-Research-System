from __future__ import annotations

import hashlib
import json
from pathlib import Path
from xml.etree import ElementTree
import zipfile
import os
import shutil
import subprocess
import sys

import pytest
from docx import Document
from openpyxl import load_workbook
from app.reporting import generate_report_bundle
from app.reporting.bundle import portable_image_links
from app.reporting.generators import write_research_deck, write_research_docx, write_results_workbook
from app.storage.run_store import RunStore
from app.bridge.report_service import download_path, export_bundle
from app.storage.report_skill_store import import_report_skill, report_skill_options, select_report_skills, selected_report_skills
from app.harness.skills.registry import load_selected_skills
from app.storage.artifact_store import ArtifactStore
from app.storage.run_store import RunHandle


def _approved_conversion_run(tmp_path: Path) -> RunHandle:
    run = RunStore(tmp_path).create(task="on-demand document conversion", project="synthetic_regression")
    reference = ArtifactStore(run).write_metadata(metadata={"schema": "report.v1", "agent": "writing",
        "project": run.project, "deliverable_type": "research_report", "target_audience": "research team",
        "chain_refs": {}}, body="# 文件转换测试\n\n此文档用于验证文件转换，不作研究结果声明。", expected_schema="report.v1")
    ArtifactStore(run).approve(reference)
    return run


def test_approval_default_saves_portable_markdown_and_images_without_office(tmp_path: Path) -> None:
    from PIL import Image
    run = _approved_conversion_run(tmp_path)
    report = run.root / "writing/research_report.approved.md"
    report.write_text(report.read_text() + '\n![正文图片](execution/conversion-fixture.png "实验图片")\n')
    original = report.read_bytes()
    plot = run.root / "execution/conversion-fixture.png"
    Image.new("RGB", (16, 16), "white").save(plot)
    image_bytes = plot.read_bytes()
    bundle = generate_report_bundle(run, actor="inline_approve")
    assert bundle["metadata"]["requested_formats"] == []
    assert [item["kind"] for item in bundle["metadata"]["deliverables"]] == ["markdown"]
    assert not list(run.root.rglob("*.xlsx"))
    assert not list(run.root.rglob("*.docx"))
    assert not list(run.root.rglob("*.pptx"))
    assert (run.root / "writing/research_report.approved.md").read_bytes() == original
    md = run.root / bundle["metadata"]["deliverables"][0]["path"]
    image = bundle["metadata"]["images"][0]
    saved = run.root / image["path"]
    assert saved.read_bytes() == image_bytes
    assert f"(images/{saved.name})" in md.read_text(encoding="utf-8")
    assert f'![正文图片](images/{saved.name} "实验图片")' in md.read_text(encoding="utf-8")
    assert '](execution/conversion-fixture.png' not in md.read_text(encoding="utf-8")
    package = download_path(run, "report_materials.zip", bundle["manifest"])
    with zipfile.ZipFile(package) as archive:
        assert archive.read("research_report.md").decode() == md.read_text(encoding="utf-8")
        assert archive.read(f"images/{saved.name}") == image_bytes
        assert archive.testzip() is None
    plot.write_bytes(b"source changed after saving")
    assert download_path(run, saved.name, bundle["manifest"]).read_bytes() == image_bytes


def test_portable_links_preserve_unbundled_sources_and_code_examples() -> None:
    images = [{'source_path': 'execution/关键 图片.png', 'path': 'writing/deliverables/id/images/hash-关键 图片.png'}]
    source = '![已保存](<../execution/关键 图片.png> "标题")\n![未保存](execution/other.png)\n![外部](https://example.com/a.png)\n```md\n![示例](<execution/关键 图片.png>)\n```\n'
    result = portable_image_links(source, images)
    assert '![已保存](images/hash-%E5%85%B3%E9%94%AE%20%E5%9B%BE%E7%89%87.png "标题")' in result
    assert '![未保存](execution/other.png)' in result
    assert '![外部](https://example.com/a.png)' in result
    assert '```md\n![示例](<execution/关键 图片.png>)\n```' in result


def test_formats_generate_individually_and_preserve_other_current_files(tmp_path: Path) -> None:
    run = _approved_conversion_run(tmp_path)
    excel = export_bundle(run, formats=("excel",))
    assert list(run.root.rglob("*.xlsx")) and not list(run.root.rglob("*.docx")) and not list(run.root.rglob("*.pptx"))
    saved_excel = next(item for item in excel["metadata"]["deliverables"] if item["kind"] == "excel")
    word = export_bundle(run, formats=("word",))
    assert not list(run.root.rglob("*.pptx"))
    assert next(item for item in word["metadata"]["deliverables"] if item["kind"] == "excel") == saved_excel
    slides = export_bundle(run, formats=("powerpoint",))
    assert {item["kind"] for item in slides["metadata"]["deliverables"]} == {"markdown", "excel", "word", "powerpoint"}
    assert download_path(run, Path(saved_excel["path"]).name, slides["manifest"]).is_file()
    # Changing recorded inputs invalidates re-use, without deleting old exports.
    (run.root / "execution/metrics.json").write_text("[]", encoding="utf-8")
    updated = export_bundle(run, formats=("word",))
    assert {item["kind"] for item in updated["metadata"]["deliverables"]} == {"markdown", "word"}
    assert (run.root / saved_excel["path"]).is_file()


def test_unknown_office_format_is_rejected_before_writing(tmp_path: Path) -> None:
    run = _approved_conversion_run(tmp_path)
    with pytest.raises(ValueError, match="仅支持"):
        export_bundle(run, formats=("unknown",))
    assert not (run.root / "writing/deliverables").exists()


def test_editable_office_formats_keep_full_text_and_numeric_precision(tmp_path: Path) -> None:
    body = "# 完整报告\n\n" + "正文内容，不作实验断言。\n\n" * 200 + "## 最后结论\n\n末尾完整保留。\n\n| 条件 | 值 |\n| --- | --- |\n| seed | 2026 |\n"
    pack = {"task": "conversion fixture", "run_id": "file_fixture", "report_markdown": body,
        "metrics": [{"experiment_id": "numeric_fixture", "RES": 23.759658938026675, "loss": .000001234}],
        "source_refs": ["execution/metrics.json"], "curves": [{"experiment_id": "numeric_fixture", "points": [
            {"optimizer_step": 1, "training_loss": .002, "lr": .0004}, {"optimizer_step": 2, "training_loss": .001, "lr": .0004}]}]}
    for name, writer in (("report.docx", write_research_docx), ("report.xlsx", write_results_workbook), ("report.pptx", write_research_deck)):
        writer(tmp_path / name, pack)
    doc = Document(tmp_path / "report.docx")
    assert "末尾完整保留。" in '\n'.join(p.text for p in doc.paragraphs)
    assert len(doc.tables) == 1
    workbook = load_workbook(tmp_path / "report.xlsx")
    assert workbook["Metrics"]["B2"].value == pytest.approx(23.759658938026675, rel=0, abs=1e-14)
    assert workbook["Metrics"]["C2"].value == .000001234
    assert workbook["Metrics"].freeze_panes == "A2"
    assert len(workbook["Steps_1"]._charts) == 1
    with zipfile.ZipFile(tmp_path / "report.pptx") as archive:
        slides = [name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
        assert len(slides) > 3
        assert "末尾完整保留。" in ''.join(archive.read(name).decode() for name in slides)
        assert any("<a:tbl>" in archive.read(name).decode() for name in slides)
        assert any('uri="http://schemas.openxmlformats.org/drawingml/2006/table"' in archive.read(name).decode() for name in slides)
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                ElementTree.fromstring(archive.read(name))


def test_export_requires_approval_and_rejects_tampered_or_unlisted_download(tmp_path: Path) -> None:
    run = RunStore(tmp_path).create(task="document conversion", project="synthetic_regression")
    with pytest.raises(ValueError, match="批准"):
        export_bundle(run)
    reference = ArtifactStore(run).write_metadata(metadata={"schema": "report.v1", "agent": "writing",
        "project": run.project, "deliverable_type": "research_report", "target_audience": "research team",
        "chain_refs": {}}, body="# 文件转换测试\n\n不是研究成功声明。", expected_schema="report.v1")
    ArtifactStore(run).approve(reference)
    bundle = generate_report_bundle(run, formats=("word",))
    assert bundle["metadata"]["generator"] == "office_editable"
    path = download_path(run, "research_report.docx", bundle["manifest"])
    path.write_bytes(path.read_bytes() + b"tampered")
    with pytest.raises(ValueError, match="完整性"):
        download_path(run, path.name, bundle["manifest"])
    with pytest.raises(FileNotFoundError):
        download_path(run, "report_data_pack.v1.json", bundle["manifest"])
    with pytest.raises(ValueError):
        download_path(run, "../secret")
    next_bundle = generate_report_bundle(run, formats=("word",))
    assert next_bundle["metadata"]["data_pack"] != bundle["metadata"]["data_pack"]


def test_skill_import_versions_permissions_and_actual_selected_context(tmp_path: Path) -> None:
    content = "---\nname: team-report\ndescription: 团队报告规范\n---\n区分事实和假设，完整列出局限。"
    imported = import_report_skill(content, root=tmp_path)
    chosen = select_report_skills("team", [imported["id"]], [], root=tmp_path)
    ids = selected_report_skills("team", root=tmp_path)
    assert chosen["selected"] == ids == [f"{imported['id']}@{imported['version']}"]
    selection = load_selected_skills(ids, granted_tools=[], project="team", registry_path=tmp_path / "configs/skills.yaml")
    assert "区分事实和假设" in selection.context
    assert selection.manifest["skills"][0]["content_sha256"] == hashlib.sha256(content.encode()).hexdigest()
    changed = import_report_skill(content + "\n不宣称统计显著性。", root=tmp_path)
    assert changed["id"] != imported["id"]
    assert selected_report_skills("team", root=tmp_path) == ids
    unavailable = import_report_skill(content.replace("description:", "required_tools: [not.granted]\ndescription:"), root=tmp_path)
    option = next(item for item in report_skill_options("team", [], root=tmp_path)["options"] if item["id"] == unavailable["id"])
    assert not option["available"]
    with pytest.raises(ValueError, match="ungranted"):
        select_report_skills("team", [unavailable["id"]], [], root=tmp_path)
    with pytest.raises(ValueError):
        import_report_skill("正文没有 name 和 description", root=tmp_path)


def test_spreadsheet_text_cannot_become_formula(tmp_path: Path) -> None:
    path = tmp_path / "text.xlsx"
    write_results_workbook(path, {"report_markdown": "=HYPERLINK(\"https://invalid.example\")"})
    book = load_workbook(path)
    assert book["Report"]["A2"].data_type == "s"


def test_long_slide_heading_keeps_text_without_duplicate_shapes(tmp_path: Path) -> None:
    heading = "很长的证据章节标题" * 24
    path = tmp_path / "long.pptx"
    write_research_deck(path, {"report_markdown": f"# {heading}\n\n完整正文仍须保留。"})
    texts = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                node = ElementTree.fromstring(archive.read(name))
                ids = [item.attrib["id"] for item in node.findall(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")]
                assert len(ids) == len(set(ids))
                texts.extend(item.text or "" for item in node.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}t"))
    combined = ''.join(texts)
    assert heading in combined
    assert "完整正文仍须保留。" in combined


def test_malformed_skill_preferences_are_rejected_without_overwriting(tmp_path: Path) -> None:
    path = tmp_path / "configs/report_skill_selection.yaml"
    path.parent.mkdir()
    path.write_text("- invalid\n", encoding="utf-8")
    with pytest.raises(ValueError):
        selected_report_skills("team", root=tmp_path)
    with pytest.raises(ValueError):
        select_report_skills("team", [], [], root=tmp_path)
    assert path.read_text(encoding="utf-8") == "- invalid\n"


def test_imported_skill_enters_real_writing_messages_and_resume_keeps_version(tmp_path: Path) -> None:
    from app.settings import repo_root
    root = tmp_path / "runtime"
    shutil.copytree(repo_root() / "configs", root / "configs", ignore=shutil.ignore_patterns("skills", "report_skill_selection.yaml", "*.lock"))
    shutil.copytree(repo_root() / "templates", root / "templates")
    shutil.copytree(repo_root() / "backend/app/agents/writing", root / "backend/app/agents/writing")
    shutil.copytree(repo_root() / "backend/app/harness/schema/schemas", root / "backend/app/harness/schema/schemas")
    (root / "configs/skills.yaml").write_text("version: 1\nskills: {}\n")
    content = "---\nname: actual-writing-context\ndescription: Context-only verification\n---\nREPORT_SKILL_CONTEXT_MARKER: separate observations from hypotheses."
    imported = import_report_skill(content, root=root)
    select_report_skills("generic", [imported["id"]], [], root=root)
    script = '''
import asyncio, json
from pathlib import Path
from app.agents.writing.agent import WritingAgent
from app.agents.base import RunRequest
from app.storage.run_store import RunStore
from app.bridge.report_skill_binding import frozen_report_skills
from app.storage.report_skill_store import select_report_skills
async def check():
    run = RunStore(Path("TEST_RUN_ROOT")).create(task="skill-context", project="generic")
    ids = frozen_report_skills(run, "writing", create=True)
    assert ids
    select_report_skills("generic", [], [])
    assert frozen_report_skills(run, "writing", create=True) == ids
    agent = WritingAgent()
    request = RunRequest(project="generic", user_request="Context inspection only; do not call a model.",
        extra={"run_root":str(run.root), "invocation_id":"contextmarker", "skills":ids})
    context = await agent.build_context(request)
    messages = agent._messages_for_context(request, context, purpose="loop")
    assert any("REPORT_SKILL_CONTEXT_MARKER" in (m.content or "") for m in messages)
    again = await agent.build_context(request)
    assert again.metadata["skills"] == context.metadata["skills"]
asyncio.run(check())
'''.replace("TEST_RUN_ROOT", str(tmp_path / "runs"))
    result = subprocess.run([sys.executable, "-c", script], env={**os.environ, "MARS_RUNTIME_ROOT": str(root)}, capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr
